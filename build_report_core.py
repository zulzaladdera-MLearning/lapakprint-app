# -*- coding: utf-8 -*-
"""
Full Academic Report Generator: LapakPrint Knowledge Domain & PRD (30+ Pages)
All text in Times New Roman, comprehensive tables, in-depth academic discourse.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color="94A3B8", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def apply_tnr_run(run, size_pt=12, bold=False, italic=False, color_rgb=None):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    else:
        run.font.color.rgb = RGBColor(15, 23, 42)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Times New Roman" w:hAnsi="Times New Roman" w:cs="Times New Roman"/>')
    rPr.append(rFonts)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=0, space_after=3, line_spacing=1.12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        apply_tnr_run(r, size_pt=12)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    apply_tnr_run(r, size_pt=13.5, bold=True, color_rgb=RGBColor(15, 23, 42))
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    apply_tnr_run(r, size_pt=12, bold=True, color_rgb=RGBColor(2, 132, 199))
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    apply_tnr_run(r, size_pt=11.5, bold=True, italic=True, color_rgb=RGBColor(51, 65, 85))
    return p

def add_bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.12
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        apply_tnr_run(r_pre, size_pt=12, bold=True)
    r_body = p.add_run(text)
    apply_tnr_run(r_body, size_pt=12)
    return p

def create_table_custom(doc, headers, data, col_widths=None, header_bg="0284C7", alt_bg="F8FAFC"):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(title)
        apply_tnr_run(r, size_pt=9.5, bold=True, color_rgb=RGBColor(255, 255, 255))
        set_cell_background(hdr_cells[i], header_bg)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = col_widths[i]

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = alt_bg if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = ""
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1.5)
            p.paragraph_format.space_after = Pt(1.5)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(cell_value))
            apply_tnr_run(r, size_pt=9.0, color_rgb=RGBColor(30, 41, 59))
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=50, bottom=50, left=80, right=80)
            if col_widths and col_idx < len(col_widths):
                row_cells[col_idx].width = col_widths[col_idx]
                
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(1)
    p_spacer.paragraph_format.space_after = Pt(3)
    return table

def create_screenshot_box(doc, image_number, caption, description, image_path=None):
    import os
    if image_path and os.path.exists(image_path):
        # Hitung rasio gambar agar tidak melebihi batas halaman
        try:
            from PIL import Image
            with Image.open(image_path) as im:
                w_px, h_px = im.size
                aspect = h_px / w_px
                if aspect > 1.3:
                    # Format potret (misal screenshot mobile)
                    img_w = Inches(1.98)
                elif aspect > 0.7:
                    img_w = Inches(3.15)
                else:
                    # Format lanskap (desktop / dashboard)
                    img_w = Inches(3.93)
        except Exception:
            img_w = Inches(3.93)

        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "FFFFFF")
        set_cell_margins(cell, top=40, bottom=40, left=40, right=40)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E1"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run_img = p.add_run()
        run_img.add_picture(image_path, width=img_w)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(3)
        p_cap.paragraph_format.space_after = Pt(1)
        r_cap = p_cap.add_run(f"Gambar {image_number}: {caption}")
        apply_tnr_run(r_cap, size_pt=9.5, bold=True, color_rgb=RGBColor(30, 41, 59))
        
        if description:
            p_desc = doc.add_paragraph()
            p_desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_desc.paragraph_format.space_before = Pt(0)
            p_desc.paragraph_format.space_after = Pt(4)
            r_desc = p_desc.add_run(f"({description})")
            apply_tnr_run(r_desc, size_pt=8.5, italic=True, color_rgb=RGBColor(100, 116, 139))
    else:
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:left w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:bottom w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
                <w:right w:val="dashed" w:sz="6" w:space="0" w:color="94A3B8"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        
        r_icon = p.add_run("📷\n")
        apply_tnr_run(r_icon, size_pt=20)
        
        r_ph = p.add_run("[ TEMPELKAN TANGKAPAN LAYAR DI SINI ]\n")
        apply_tnr_run(r_ph, size_pt=11, bold=True, color_rgb=RGBColor(30, 41, 59))
        
        r_desc = p.add_run(f"({description})\n")
        apply_tnr_run(r_desc, size_pt=9.5, italic=True, color_rgb=RGBColor(100, 116, 139))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(4)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Gambar {image_number}: {caption}")
        apply_tnr_run(r_cap, size_pt=9.5, bold=True, color_rgb=RGBColor(71, 85, 105))

print("Report generator core loaded.")
